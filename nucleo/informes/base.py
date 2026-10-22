# ==============================================================================
# CENTRO DE INFORMES — Base común (Fase 2)
# ==============================================================================
"""
Patrón común para todos los informes del ERP:

- Filtros de fecha estándar (?desde=YYYY-MM-DD&hasta=YYYY-MM-DD).
- Exportación Excel con ?export=xlsx (reutiliza openpyxl, ya está en requirements).
- Contrato de respuesta estándar para que el frontend (informes.js) sea genérico:

    {
        "titulo": "Valorización de Inventario",
        "columnas": [{"campo": "codigo", "titulo": "Código"}, ...],
        "resultados": [{"codigo": "MORT-001", ...}, ...],
        "resumen": {"Valor a costo": 12847.30, ...}
    }

Todos los informes son de solo lectura y usan IsGerenteOrAdmin
(consistente con el resto del módulo de inventario/reportes gerenciales).
"""

from datetime import datetime, time, timedelta
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.response import Response
from rest_framework.views import APIView

from ..permissions import IsGerenteOrAdmin
from rest_framework.permissions import IsAuthenticated


def parse_fecha(valor):
    """'YYYY-MM-DD' -> date. Devuelve None si no viene. Lanza ValueError si es inválida."""
    if not valor:
        return None
    return datetime.strptime(str(valor).strip(), '%Y-%m-%d').date()


def rango_fechas(request, por_defecto_dias=None):
    """
    Lee ?desde= y ?hasta= y devuelve (desde_dt, hasta_dt) como datetimes
    conscientes con el rango [desde 00:00, hasta+1d 00:00).

    - Si no vienen fechas y `por_defecto_dias` está definido, usa los últimos N días.
    - Si no vienen fechas y `por_defecto_dias` es None, devuelve (None, None) = sin filtro.
    Lanza ValueError si el formato es inválido.
    """
    desde = parse_fecha(request.query_params.get('desde'))
    hasta = parse_fecha(request.query_params.get('hasta'))

    if por_defecto_dias is not None and not desde and not hasta:
        hoy = timezone.localdate()
        desde = hoy - timedelta(days=por_defecto_dias)
        hasta = hoy

    desde_dt = timezone.make_aware(datetime.combine(desde, time.min)) if desde else None
    hasta_dt = timezone.make_aware(datetime.combine(hasta + timedelta(days=1), time.min)) if hasta else None
    return desde_dt, hasta_dt


class InformeBaseAPIView(APIView):
    """Base de los informes: permisos de gerencia + exportación Excel."""
    permission_classes = [IsAuthenticated, IsGerenteOrAdmin]
    exportable = True

    def exportar_excel(self, request, titulo, columnas, resultados, filename, resumen=None):
        """
        Si la petición trae ?export=xlsx devuelve un HttpResponse con el Excel.
        Si no, devuelve None (el caller responde JSON normal).
        """
        if request.query_params.get('export') != 'xlsx':
            return None

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = titulo[:31] or 'Informe'

        # Título
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columnas), 1))
        celda_titulo = ws.cell(row=1, column=1, value=titulo)
        celda_titulo.font = Font(bold=True, size=14)

        # Encabezados
        borde = Border(*[Side(style='thin', color='CCCCCC')] * 4)
        for j, col in enumerate(columnas, start=1):
            celda = ws.cell(row=2, column=j, value=col.get('titulo', col.get('campo', '')))
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='343A40')
            celda.alignment = Alignment(horizontal='center')
            celda.border = borde
            ws.column_dimensions[celda.column_letter].width = max(12, len(str(col.get('titulo', ''))) + 4)

        # Datos
        for i, fila in enumerate(resultados, start=3):
            for j, col in enumerate(columnas, start=1):
                valor = fila.get(col.get('campo'))
                celda = ws.cell(row=i, column=j, value=valor)
                celda.border = borde
                if col.get('formato') in ('moneda', 'numero'):
                    celda.number_format = '#,##0.00'
                elif col.get('formato') == 'entero':
                    celda.number_format = '#,##0'

        # Resumen al pie
        if resumen:
            fila_res = len(resultados) + 4
            ws.cell(row=fila_res, column=1, value='RESUMEN').font = Font(bold=True)
            for k, (clave, valor) in enumerate(resumen.items(), start=1):
                ws.cell(row=fila_res + k, column=1, value=str(clave)).font = Font(bold=True)
                celda = ws.cell(row=fila_res + k, column=2, value=valor)
                if isinstance(valor, (int, float)):
                    celda.number_format = '#,##0.00'

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    def respuesta_invalida(self, mensaje):
        return Response({"error": mensaje}, status=status.HTTP_400_BAD_REQUEST)
